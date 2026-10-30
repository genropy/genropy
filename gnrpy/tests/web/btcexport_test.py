"""Regression test for issue #1035: gridcall() must not crash on an empty grid.

When the client-side grid selection is empty, genro_grid.js's
mixin_serverAction sends data=None to the server (storebag() on an empty
grid yields nothing), while struct is always sent. BaseResourceExport.
gridcall() used to dereference data.output('grid') unconditionally,
raising AttributeError: 'NoneType' object has no attribute 'output'.
"""

import os
import shutil
import tempfile

import pytest
import openpyxl

from gnr.app.gnrapp import GnrApp
from core.common import BaseGnrTest
from gnr.core.gnrbag import Bag
from gnr.web.batch.btcexport import BaseResourceExport


def setup_module(module):
    BaseGnrTest.setup_class()


def teardown_module(module):
    BaseGnrTest.teardown_class()


class _FakeStorageNode(str):
    """Minimal stand-in for a genropy StorageNode: a real filesystem path
    string with the .url() accessor gridcall()'s post_process() expects.
    """

    def url(self):
        return str(self)


class _FakeSite:
    """Provides just the slice of Site used by gridcall(): a storageNode()
    that does real file I/O against a temporary directory, without booting
    a full GnrWsgiSite/daemon (not exercised by the code path under test).
    """

    def __init__(self, base_dir):
        self.base_dir = base_dir

    def storageNode(self, storage, *parts, **kwargs):
        path = os.path.join(self.base_dir, storage.replace(':', '_'), *parts)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        return _FakeStorageNode(path)


class _FakePage:
    """Narrowest real stand-in for a webpage. BaseResourceBatch.__init__()
    only needs .db, .btc and .getUuid(); the gridcall() export path
    additionally needs .locale and .site.storageNode(). Everything else
    exercised by the test (self.db, prepareExportCols, rowFromAttr, the
    writer, the Bag) is the genuine implementation.
    """

    def __init__(self, db, output_dir):
        self.db = db
        self.btc = object()
        self.locale = 'en'
        self.site = _FakeSite(output_dir)

    def getUuid(self):
        return 'test-uuid-0001'


@pytest.fixture(scope='module')
def sqlite_temp_dir():
    tmpdir = tempfile.mkdtemp()
    try:
        yield tmpdir
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.fixture(scope='module')
def sqlite_db(sqlite_temp_dir):
    app = GnrApp('test_invoice', db_attrs=dict(
        implementation='sqlite',
        dbname=os.path.join(sqlite_temp_dir, 'testing'),
    ))
    app.db.model.check(applyChanges=True)
    return app.db


def _build_struct():
    """A real struct Bag shaped as _prepareExportCols_struct() expects: it
    pops 'info', then walks views -> rows -> cells reading the field/name/
    dtype/hidden/columnset/caption_field/group_aggr cell attributes.
    """
    struct = Bag()
    row = Bag()
    row.setItem('description', None, field='description', name='Description', dtype='A')
    row.setItem('amount', None, field='amount', name='Amount', dtype='N')
    view = Bag()
    view.setItem('row_0', row)
    struct.setItem('view_0', view)
    return struct


def test_gridcall_empty_grid_exports_headers_only(sqlite_db, tmp_path):
    """An empty grid (data=None) must produce a headers-only file instead
    of raising AttributeError."""
    page = _FakePage(sqlite_db, str(tmp_path))
    resource = BaseResourceExport(page=page, resource_table=None)

    fileurl = resource.gridcall(data=None, struct=_build_struct(),
                                 export_mode='xls', filename='empty_grid_export')

    assert fileurl
    assert os.path.isfile(fileurl)

    workbook = openpyxl.load_workbook(fileurl)
    sheet = workbook.active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == ['Description', 'Amount']


def _grouped_struct():
    """A grouped view struct: one group column, a sum and the row counter,
    named as the grouped view's select method names its columns."""
    struct = Bag()
    row = Bag()
    row.setItem('product_type_id', None, field='product_type_id', name='Type', dtype='A')
    row.setItem('unit_price', None, field='unit_price', name='Price', dtype='N', group_aggr='sum')
    row.setItem('_grp_count', None, field='_grp_count', name='Count', dtype='L', group_aggr='sum')
    view = Bag()
    view.setItem('row_0', row)
    struct.setItem('view_0', view)
    return struct


def test_gridcall_exports_a_frozen_grouped_selection_as_the_uploaded_rows(sqlite_db, tmp_path):
    """#1773: a grouped view exports from its frozen selection instead of
    the rows the browser uploaded. A GROUP BY selection has no pkey column,
    and the file must be the one the uploaded rows produced."""
    product = sqlite_db.table('invc.product')
    for code, product_type_id, unit_price in (('P1', 'T1', 10), ('P2', 'T1', 15),
                                              ('P3', 'T2', 7), ('P4', 'T3', 30)):
        product.insert(dict(code=code, description=code, product_type_id=product_type_id,
                            unit_price=unit_price))
    grouped = product.query(columns='$product_type_id,sum($unit_price) AS unit_price_sum,'
                                    'count(*) AS _grp_count_sum',
                            group_by='$product_type_id', order_by='$product_type_id').selection()
    uploaded = grouped.output('grid')
    grouped.setKey('rowidx')
    freeze_path = str(tmp_path / 'grouped')
    grouped.freeze(freeze_path, freezePkeys=True)
    frozen = product.frozenSelection(freeze_path)
    page = _FakePage(sqlite_db, str(tmp_path))

    def exported(data, filename):
        resource = BaseResourceExport(page=page, resource_table=None)
        fileurl = resource.gridcall(data=data, struct=_grouped_struct(),
                                    export_mode='xls', filename=filename)
        return [[cell.value for cell in row] for row in openpyxl.load_workbook(fileurl).active.iter_rows()]

    from_selection = exported(frozen, 'grouped_from_selection')
    assert from_selection == exported(uploaded, 'grouped_from_upload')
    assert from_selection == [['Type', 'Price', 'Count'], ['T1', 25, 2], ['T2', 7, 1], ['T3', 30, 1]]
